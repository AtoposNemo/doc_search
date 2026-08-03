import os
import sys
import csv
import argparse


# === Supported file extensions ===
SUPPORTED_EXTENSIONS = {'.doc', '.docx', '.txt', '.csv', '.xlsx', '.xls', '.pdf'}

# Extensions that can be read with pure Python (parallel, fast)
_PURE_PYTHON_EXTS = {'.docx', '.txt', '.csv', '.xlsx', '.pdf'}
# Extensions that require COM automation (sequential, slower)
_COM_EXTS = {'.doc', '.xls'}


class NoOfficeSoftwareError(Exception):
    """Raised when neither Microsoft Word nor WPS Office is installed
    but .doc files need to be read."""
    pass


class NoExcelSoftwareError(Exception):
    """Raised when neither Microsoft Excel nor WPS Office is installed
    but .xls files need to be read."""
    pass


# COM ProgIDs to try, in order of preference.
# Word.Application  — Microsoft Word (primary)
# KWPS.Application  — WPS Office (older versions)
# wps.Application   — WPS Office (newer versions)
_OFFICE_PROG_IDS = ('Word.Application', 'KWPS.Application', 'wps.Application')

# Excel COM ProgIDs to try, in order of preference.
# Excel.Application — Microsoft Excel (primary)
# Ket.Application   — WPS Spreadsheet (older versions)
# et.Application    — WPS Spreadsheet (newer versions)
_EXCEL_PROG_IDS = ('Excel.Application', 'Ket.Application', 'et.Application')


def _create_office_app():
    """Create a COM office application instance.

    Tries Microsoft Word first, then falls back to WPS Office.
    Returns the application object configured for fast, hidden, read-only use.
    Raises NoOfficeSoftwareError if neither Word nor WPS is installed.
    """
    import win32com.client

    word_app = None
    last_error = None
    for prog_id in _OFFICE_PROG_IDS:
        try:
            word_app = win32com.client.DispatchEx(prog_id)
            break
        except Exception as e:
            last_error = e
            continue

    if word_app is None:
        raise NoOfficeSoftwareError(
            "未检测到 Microsoft Word 或 WPS Office。\n"
            "搜索 .doc 格式文件需要安装其中之一。\n"
            "如仅搜索 .docx 文件则无需安装。"
        )

    # Configure for fast, hidden, non-intrusive operation.
    # All settings are wrapped in try/except because WPS may not support
    # some Word-specific properties.
    try:
        word_app.Visible = False
    except Exception:
        pass
    try:
        word_app.DisplayAlerts = 0  # wdAlertsNone
    except Exception:
        pass
    try:
        word_app.ScreenUpdating = False  # big speedup: no UI redraw
    except Exception:
        pass
    try:
        word_app.AutomationSecurity = 3  # disable macros
    except Exception:
        pass
    try:
        word_app.NormalTemplate.Saved = True  # skip "save Normal?" prompt
    except Exception:
        pass

    return word_app


def _create_excel_app():
    """Create a COM Excel application instance.

    Tries Microsoft Excel first, then falls back to WPS Spreadsheet.
    Returns the application object configured for fast, hidden, read-only use.
    Raises NoExcelSoftwareError if neither Excel nor WPS is installed.
    """
    import win32com.client

    excel_app = None
    for prog_id in _EXCEL_PROG_IDS:
        try:
            excel_app = win32com.client.DispatchEx(prog_id)
            break
        except Exception:
            continue

    if excel_app is None:
        raise NoExcelSoftwareError(
            "未检测到 Microsoft Excel 或 WPS Office。\n"
            "搜索 .xls 格式文件需要安装其中之一。\n"
            "如仅搜索 .xlsx 文件则无需安装。"
        )

    try:
        excel_app.Visible = False
    except Exception:
        pass
    try:
        excel_app.DisplayAlerts = False
    except Exception:
        pass
    try:
        excel_app.ScreenUpdating = False
    except Exception:
        pass

    return excel_app


# === Text encoding helper ===
def _read_text_file(file_path):
    """Read a text file trying common encodings (utf-8, gbk, latin-1)."""
    for encoding in ('utf-8-sig', 'utf-8', 'gbk', 'gb18030', 'latin-1'):
        try:
            with open(file_path, 'r', encoding=encoding) as f:
                return f.read()
        except (UnicodeDecodeError, LookupError):
            continue
    # latin-1 should never fail, but fallback just in case
    with open(file_path, 'r', encoding='latin-1', errors='replace') as f:
        return f.read()


def read_docx(file_path):
    from docx import Document
    try:
        doc = Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text.append(cell.text)
        return '\n'.join(full_text)
    except Exception as e:
        print(f"Warning: Failed to read {file_path}: {str(e)}", file=sys.stderr)
        return ""


def read_doc(file_path, word_app):
    try:
        doc = word_app.Documents.Open(
            FileName=os.path.abspath(file_path),
            ReadOnly=True,
            AddToRecentFiles=False,
            Visible=False
        )
        text = doc.Content.Text
        doc.Close(SaveChanges=0)  # wdDoNotSaveChanges
        return text
    except Exception as e:
        print(f"Warning: Failed to read {file_path}: {str(e)}", file=sys.stderr)
        return ""


def read_txt(file_path):
    """Read a plain text file (.txt) with encoding auto-detection."""
    try:
        return _read_text_file(file_path)
    except Exception as e:
        print(f"Warning: Failed to read {file_path}: {str(e)}", file=sys.stderr)
        return ""


def read_csv(file_path):
    """Read a CSV file and return all cell text joined by newlines."""
    try:
        rows = []
        for encoding in ('utf-8-sig', 'utf-8', 'gbk', 'gb18030', 'latin-1'):
            try:
                with open(file_path, 'r', encoding=encoding, newline='') as f:
                    reader = csv.reader(f)
                    for row in reader:
                        rows.append('\t'.join(row))
                return '\n'.join(rows)
            except (UnicodeDecodeError, LookupError):
                continue
        return '\n'.join(rows)
    except Exception as e:
        print(f"Warning: Failed to read {file_path}: {str(e)}", file=sys.stderr)
        return ""


def read_xlsx(file_path):
    """Read an Excel .xlsx file using openpyxl."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(file_path, read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append('\t'.join(cells))
        wb.close()
        return '\n'.join(lines)
    except Exception as e:
        print(f"Warning: Failed to read {file_path}: {str(e)}", file=sys.stderr)
        return ""


def read_xls(file_path, excel_app):
    """Read an Excel .xls file using COM automation (Excel or WPS)."""
    try:
        wb = excel_app.Workbooks.Open(
            FileName=os.path.abspath(file_path),
            ReadOnly=True,
            AddToRecentFiles=False
        )
        lines = []
        for ws in wb.Worksheets:
            used = ws.UsedRange
            rows = used.Rows.Count
            cols = used.Columns.Count
            for r in range(1, rows + 1):
                cells = []
                for c in range(1, cols + 1):
                    val = ws.Cells(r, c).Value
                    if val is not None:
                        cells.append(str(val))
                if cells:
                    lines.append('\t'.join(cells))
        wb.Close(SaveChanges=False)
        return '\n'.join(lines)
    except Exception as e:
        print(f"Warning: Failed to read {file_path}: {str(e)}", file=sys.stderr)
        return ""


def read_pdf(file_path):
    """Read a PDF file using PyPDF2."""
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(file_path)
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
        return '\n'.join(pages)
    except Exception as e:
        print(f"Warning: Failed to read {file_path}: {str(e)}", file=sys.stderr)
        return ""


def extract_text(file_path, word_app=None, excel_app=None):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".docx":
        return read_docx(file_path)
    elif ext == ".doc":
        return read_doc(file_path, word_app)
    elif ext == ".txt":
        return read_txt(file_path)
    elif ext == ".csv":
        return read_csv(file_path)
    elif ext == ".xlsx":
        return read_xlsx(file_path)
    elif ext == ".xls":
        return read_xls(file_path, excel_app)
    elif ext == ".pdf":
        return read_pdf(file_path)
    return ""


def get_paragraphs_docx(file_path):
    from docx import Document
    try:
        doc = Document(file_path)
        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text:
                        paragraphs.append(text)
        return paragraphs
    except Exception as e:
        print(f"Warning: Failed to read {file_path}: {str(e)}", file=sys.stderr)
        return []


def get_paragraphs_doc(file_path, word_app):
    try:
        doc = word_app.Documents.Open(
            FileName=os.path.abspath(file_path),
            ReadOnly=True,
            AddToRecentFiles=False,
            Visible=False
        )
        paragraphs = []
        for para in doc.Paragraphs:
            text = para.Range.Text.strip()
            if text:
                paragraphs.append(text)
        doc.Close(SaveChanges=0)
        return paragraphs
    except Exception as e:
        print(f"Warning: Failed to read {file_path}: {str(e)}", file=sys.stderr)
        return []


def get_paragraphs_txt(file_path):
    """Extract paragraphs from a .txt file (split by newlines)."""
    try:
        text = _read_text_file(file_path)
        return [line.strip() for line in text.split('\n') if line.strip()]
    except Exception as e:
        print(f"Warning: Failed to read {file_path}: {str(e)}", file=sys.stderr)
        return []


def get_paragraphs_csv(file_path):
    """Extract paragraphs from a .csv file (each row becomes a paragraph)."""
    try:
        paragraphs = []
        for encoding in ('utf-8-sig', 'utf-8', 'gbk', 'gb18030', 'latin-1'):
            try:
                with open(file_path, 'r', encoding=encoding, newline='') as f:
                    reader = csv.reader(f)
                    for row in reader:
                        text = '\t'.join(row).strip()
                        if text:
                            paragraphs.append(text)
                return paragraphs
            except (UnicodeDecodeError, LookupError):
                continue
        return paragraphs
    except Exception as e:
        print(f"Warning: Failed to read {file_path}: {str(e)}", file=sys.stderr)
        return []


def get_paragraphs_xlsx(file_path):
    """Extract paragraphs from an .xlsx file (each row becomes a paragraph)."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(file_path, read_only=True, data_only=True)
        paragraphs = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                text = '\t'.join(cells).strip()
                if text:
                    paragraphs.append(text)
        wb.close()
        return paragraphs
    except Exception as e:
        print(f"Warning: Failed to read {file_path}: {str(e)}", file=sys.stderr)
        return []


def get_paragraphs_xls(file_path, excel_app):
    """Extract paragraphs from an .xls file using COM (each row = paragraph)."""
    try:
        wb = excel_app.Workbooks.Open(
            FileName=os.path.abspath(file_path),
            ReadOnly=True,
            AddToRecentFiles=False
        )
        paragraphs = []
        for ws in wb.Worksheets:
            used = ws.UsedRange
            rows = used.Rows.Count
            cols = used.Columns.Count
            for r in range(1, rows + 1):
                cells = []
                for c in range(1, cols + 1):
                    val = ws.Cells(r, c).Value
                    if val is not None:
                        cells.append(str(val))
                text = '\t'.join(cells).strip()
                if text:
                    paragraphs.append(text)
        wb.Close(SaveChanges=False)
        return paragraphs
    except Exception as e:
        print(f"Warning: Failed to read {file_path}: {str(e)}", file=sys.stderr)
        return []


def get_paragraphs_pdf(file_path):
    """Extract paragraphs from a PDF file (split by newlines per page)."""
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(file_path)
        paragraphs = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                for line in text.split('\n'):
                    line = line.strip()
                    if line:
                        paragraphs.append(line)
        return paragraphs
    except Exception as e:
        print(f"Warning: Failed to read {file_path}: {str(e)}", file=sys.stderr)
        return []


def is_english(text):
    english_chars = sum(1 for c in text if c.isalpha() and ord(c) < 128)
    total_chars = sum(1 for c in text if c.isalpha())
    if total_chars == 0:
        return False
    return english_chars / total_chars > 0.5


def truncate_text(text, max_length=300):
    if len(text) <= max_length:
        return text
    return text[:max_length] + '...'


def _build_translation(text, search_str, search_len, match_func, find_func, context_chars):
    """Build a translation excerpt, centering on the keyword if present."""
    if match_func(text, search_str):
        t_idx = find_func(text, search_str)
        t_match_end = t_idx + search_len
        t_start = max(0, t_idx - context_chars)
        t_end = min(len(text), t_match_end + context_chars)
        t_prefix = '...' if t_start > 0 else ''
        t_suffix = '...' if t_end < len(text) else ''
        excerpt = t_prefix + text[t_start:t_end] + t_suffix
        ms = len(t_prefix) + (t_idx - t_start)
        me = ms + search_len
        return excerpt, ms, me
    else:
        return truncate_text(text, 400), -1, -1


def get_match_excerpts(file_path, search_str, case_sensitive=False, context_chars=200, word_app=None, excel_app=None):
    """Extract excerpts centered around the search keyword.

    - Keyword always at the center of the excerpt (context_chars before/after)
    - Searches ALL paragraphs (both Chinese and English)
    - Bilingual docs: ZH paragraphs listed first, then EN paragraphs in same order.
      So zh_paragraphs[i] corresponds to en_paragraphs[i].
    - Returns translation_match_start/end for translation highlighting
    """
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".docx":
        paragraphs = get_paragraphs_docx(file_path)
    elif ext == ".doc":
        paragraphs = get_paragraphs_doc(file_path, word_app)
    elif ext == ".txt":
        paragraphs = get_paragraphs_txt(file_path)
    elif ext == ".csv":
        paragraphs = get_paragraphs_csv(file_path)
    elif ext == ".xlsx":
        paragraphs = get_paragraphs_xlsx(file_path)
    elif ext == ".xls":
        paragraphs = get_paragraphs_xls(file_path, excel_app)
    elif ext == ".pdf":
        paragraphs = get_paragraphs_pdf(file_path)
    else:
        return []

    if not paragraphs or not search_str:
        return []

    # Split into zh/en lists while preserving order correspondence
    zh_paragraphs = []
    en_paragraphs = []
    for para in paragraphs:
        if is_english(para):
            en_paragraphs.append(para)
        else:
            zh_paragraphs.append(para)

    excerpts = []
    search_len = len(search_str)

    if case_sensitive:
        target_str = search_str
        match_func = lambda text, target: target in text
        find_func = lambda text, target: text.find(target)
    else:
        target_str = search_str.lower()
        match_func = lambda text, target: target in text.lower()
        find_func = lambda text, target: text.lower().find(target)

    # Helper: build excerpt from a matched paragraph
    def make_excerpt(para, idx):
        match_end_idx = idx + search_len
        excerpt_start = max(0, idx - context_chars)
        excerpt_end = min(len(para), match_end_idx + context_chars)
        prefix = '...' if excerpt_start > 0 else ''
        suffix = '...' if excerpt_end < len(para) else ''
        excerpt_text = prefix + para[excerpt_start:excerpt_end] + suffix
        ms = len(prefix) + (idx - excerpt_start)
        me = ms + search_len
        return excerpt_text, ms, me

    # Search Chinese paragraphs
    for i, zh_para in enumerate(zh_paragraphs):
        if not match_func(zh_para, target_str):
            continue

        idx = find_func(zh_para, target_str)
        excerpt_text, ms, me = make_excerpt(zh_para, idx)

        # Translation: corresponding EN paragraph by index
        translation_text = ""
        translation_match_start = -1
        translation_match_end = -1

        if i < len(en_paragraphs):
            en_para = en_paragraphs[i]
            translation_text, translation_match_start, translation_match_end = \
                _build_translation(en_para, target_str, search_len, match_func, find_func, context_chars)

        excerpts.append({
            'text': excerpt_text,
            'match_start': ms,
            'match_end': me,
            'is_english': False,
            'translation': translation_text,
            'translation_match_start': translation_match_start,
            'translation_match_end': translation_match_end,
        })

    # Search English paragraphs
    for i, en_para in enumerate(en_paragraphs):
        if not match_func(en_para, target_str):
            continue

        idx = find_func(en_para, target_str)
        excerpt_text, ms, me = make_excerpt(en_para, idx)

        # Translation: corresponding ZH paragraph by index
        translation_text = ""
        translation_match_start = -1
        translation_match_end = -1

        if i < len(zh_paragraphs):
            zh_para = zh_paragraphs[i]
            translation_text, translation_match_start, translation_match_end = \
                _build_translation(zh_para, target_str, search_len, match_func, find_func, context_chars)

        excerpts.append({
            'text': excerpt_text,
            'match_start': ms,
            'match_end': me,
            'is_english': True,
            'translation': translation_text,
            'translation_match_start': translation_match_start,
            'translation_match_end': translation_match_end,
        })

    return excerpts


def search_in_file(file_path, search_str, case_sensitive=False, word_app=None, excel_app=None):
    text = extract_text(file_path, word_app, excel_app)

    if case_sensitive:
        return search_str in text
    else:
        return search_str.lower() in text.lower()


def search_docs(folder_path, search_str, case_sensitive=False, progress_callback=None):
    results = []
    all_files = []

    for root, dirs, files in os.walk(folder_path):
        for file in files:
            ext = os.path.splitext(file)[1].lower()
            if ext in SUPPORTED_EXTENSIONS:
                all_files.append(os.path.join(root, file))

    total = len(all_files)
    print(f"Found {total} supported files to search...")

    if total == 0:
        return results

    # Separate by type: pure Python (parallel) vs COM (sequential)
    pure_python_files = [f for f in all_files
                         if os.path.splitext(f)[1].lower() in _PURE_PYTHON_EXTS]
    doc_old_files = [f for f in all_files
                     if os.path.splitext(f)[1].lower() == '.doc']
    xls_old_files = [f for f in all_files
                     if os.path.splitext(f)[1].lower() == '.xls']

    processed = 0

    def _report_progress():
        nonlocal processed
        processed += 1
        if progress_callback:
            progress_callback(processed, total)

    # --- Pure Python files: parallel search (.docx, .txt, .csv, .xlsx, .pdf) ---
    if pure_python_files:
        from concurrent.futures import ThreadPoolExecutor, as_completed

        def _search_pure(fp):
            try:
                if search_in_file(fp, search_str, case_sensitive):
                    return fp
            except Exception as e:
                print(f"Warning: Error processing {fp}: {str(e)}", file=sys.stderr)
            return None

        max_workers = min(4, os.cpu_count() or 1)
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            future_map = {executor.submit(_search_pure, f): f for f in pure_python_files}
            for future in as_completed(future_map):
                _report_progress()
                result = future.result()
                if result is not None:
                    results.append(result)

    # --- .doc: sequential search with Office COM (Word or WPS) ---
    word_app = None
    try:
        if doc_old_files:
            if sys.platform != 'win32':
                print("Warning: .doc files require Windows with Microsoft Word "
                      "or WPS Office installed.", file=sys.stderr)
                print("Skipping .doc files...", file=sys.stderr)
            else:
                word_app = _create_office_app()

            for file_path in doc_old_files:
                _report_progress()
                print(f"Processing doc: {os.path.basename(file_path)}")
                try:
                    if search_in_file(file_path, search_str, case_sensitive, word_app):
                        results.append(file_path)
                except Exception as e:
                    print(f"Warning: Error processing {file_path}: {str(e)}",
                          file=sys.stderr)
    finally:
        if word_app:
            try:
                word_app.Quit()
            except Exception:
                pass

    # --- .xls: sequential search with Excel COM (Excel or WPS) ---
    excel_app = None
    try:
        if xls_old_files:
            if sys.platform != 'win32':
                print("Warning: .xls files require Windows with Microsoft Excel "
                      "or WPS Office installed.", file=sys.stderr)
                print("Skipping .xls files...", file=sys.stderr)
            else:
                excel_app = _create_excel_app()

            for file_path in xls_old_files:
                _report_progress()
                print(f"Processing xls: {os.path.basename(file_path)}")
                try:
                    if search_in_file(file_path, search_str, case_sensitive, None, excel_app):
                        results.append(file_path)
                except Exception as e:
                    print(f"Warning: Error processing {file_path}: {str(e)}",
                          file=sys.stderr)
    finally:
        if excel_app:
            try:
                excel_app.Quit()
            except Exception:
                pass

    return results


def main():
    parser = argparse.ArgumentParser(description='Search for a string in documents (.doc, .docx, .txt, .csv, .xlsx, .xls, .pdf)')
    parser.add_argument('search_str', help='The string to search for')
    parser.add_argument('folder_path', help='The folder path to search in')
    parser.add_argument('-c', '--case-sensitive', action='store_true', help='Case sensitive search')
    
    args = parser.parse_args()
    
    if not os.path.isdir(args.folder_path):
        print(f"Error: {args.folder_path} is not a valid directory")
        sys.exit(1)
    
    results = search_docs(args.folder_path, args.search_str, args.case_sensitive)
    
    print("\n" + "="*50)
    if results:
        print(f"Found {len(results)} files containing '{args.search_str}':")
        for result in results:
            print(result)
    else:
        print(f"No files found containing '{args.search_str}'")


if __name__ == "__main__":
    main()